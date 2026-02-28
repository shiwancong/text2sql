#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
增强版知识库构建脚本

从 tables 目录下的 Excel 文件读取表结构和测试问题，生成 RAG 知识库
"""

import json
import os
import sys
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional, Any

# Excel 读取库
try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("警告: openpyxl 未安装，将跳过 Excel 文件读取")
    print("请运行: pip install openpyxl")

# 输出目录
OUTPUT_DIR = Path(__file__).parent / 'knowledge'
DDL_DIR = OUTPUT_DIR / 'ddl'
DESC_DIR = OUTPUT_DIR / 'descriptions'
EX_DIR = OUTPUT_DIR / 'examples'

# Tables 目录
TABLES_DIR = Path(__file__).parent.parent / 'tables'

# ============================================================================
# Excel 文件读取
# ============================================================================
def read_excel_sheets(file_path: Path) -> Dict[str, Any]:
    """
    读取 Excel 文件中的表结构信息

    支持 A/BC/S 级表的不同结构:

    A/BC级结构:
    - 第4行第5列 (E列): 表名
    - 第5行第5列: 描述
    - 第10行: 字段标题行
    - 第11行开始: 字段数据

    S级结构:
    - 第4行第6列 (F列): 表名
    - 第5行第6列: 描述
    - 第10行: 字段标题行
    - 第11行开始: 字段数据
    """
    if not EXCEL_AVAILABLE:
        return {}

    result = {}
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"  警告: 无法读取 {file_path.name}: {e}")
        return {}

    # 跳过"总览" sheet
    data_sheets = [s for s in wb.sheetnames if s != '总览']

    # 判断文件类型
    is_s_file = 'S级表' in file_path.name

    for sheet_name in data_sheets:
        try:
            ws = wb[sheet_name]

            # 根据文件类型确定列位置
            if is_s_file:
                table_name_col = 6  # F列
                desc_col = 6
                notes_col = 6
                dept_col = 6
                field_name_col = 6  # F列
                name_cn_col = 4     # D列 - 字段中文名
                comment_col = 7     # G列
                type_col = 8        # H列
                code_table_col = 9  # I列
                ref_table_col = 10  # J列
                ref_field_col = 11  # K列
                required_col = 2    # B列
            else:
                table_name_col = 5  # E列
                desc_col = 5
                notes_col = 5
                dept_col = 5
                field_name_col = 5  # E列
                name_cn_col = 3     # C列
                comment_col = 6     # F列
                type_col = 7        # G列
                code_table_col = 8  # H列
                ref_table_col = 9   # I列
                ref_field_col = 10  # J列
                required_col = 2    # B列

            # 提取基本信息 (第4-7行)
            table_name = None
            description = ""
            notes = ""
            department = ""

            # 第4行: 表名
            cell = ws.cell(4, table_name_col)
            if cell.value:
                table_name = str(cell.value).strip()

            # 第5行: 描述
            cell = ws.cell(5, desc_col)
            if cell.value:
                description = str(cell.value).strip()

            # 第6行: 注意事项
            cell = ws.cell(6, notes_col)
            if cell.value:
                notes = str(cell.value).strip()

            # 第7行: 数据责任部门
            cell = ws.cell(7, dept_col)
            if cell.value:
                department = str(cell.value).strip()

            # 如果没有表名，跳过
            if not table_name:
                continue

            # 提取字段信息 (第10行开始)
            columns = []
            primary_keys = []
            foreign_keys = []

            # 从第11行开始读取字段
            for row_idx in range(11, min(ws.max_row + 1, 200)):  # 限制最多200行
                # 序号
                seq_cell = ws.cell(row_idx, 1)
                if not seq_cell.value:
                    break

                # 字段中文名
                name_cn = ws.cell(row_idx, name_cn_col).value or ""

                # 字段名
                field_name = ws.cell(row_idx, field_name_col).value or ""
                if not field_name:
                    continue

                # 数据采集要求/注释
                comment = ws.cell(row_idx, comment_col).value or ""

                # 字段类型
                data_type = ws.cell(row_idx, type_col).value or ""

                # 编码表
                code_table = ws.cell(row_idx, code_table_col).value or ""

                # 关联表
                ref_table = ws.cell(row_idx, ref_table_col).value or ""

                # 关联字段
                ref_field = ws.cell(row_idx, ref_field_col).value or ""

                # 必填标记
                required = ws.cell(row_idx, required_col).value or ""
                is_required = required == '*' or str(required).strip() == '*'

                column_info = {
                    'name': str(field_name).strip(),
                    'name_cn': str(name_cn).strip(),
                    'type': str(data_type).strip(),
                    'comment': str(comment).strip(),
                    'nullable': not is_required,
                    'code_table': str(code_table).strip() if code_table else None,
                }

                columns.append(column_info)

                # 判断主键
                if 'ID' in str(field_name).upper() and '主键' in str(comment):
                    primary_keys.append(str(field_name).strip())

                # 判断外键
                if ref_table:
                    foreign_keys.append({
                        'column': str(field_name).strip(),
                        'ref_table': str(ref_table).strip(),
                        'ref_column': str(ref_field).strip() if ref_field else 'ID'
                    })

            result[table_name] = {
                'table_name': table_name,
                'comment': description,
                'notes': notes,
                'department': department,
                'business_concept': sheet_name,  # 使用 sheet 名称作为业务概念
                'columns': columns,
                'primary_keys': primary_keys,
                'foreign_keys': foreign_keys,
                'source_file': file_path.name,
                'source_sheet': sheet_name
            }

        except Exception as e:
            print(f"  警告: 读取 sheet '{sheet_name}' 时出错: {e}")
            continue

    return result


def read_all_excel_tables() -> Dict[str, Any]:
    """读取 S级表 Excel 文件中的表结构（精简模式）"""
    all_tables = {}

    if not EXCEL_AVAILABLE:
        return all_tables

    # 只读取 S级表
    excel_files = {
        'S级表.xlsx': 'S',
    }

    print("\n=== 读取 S级表结构文件（精简模式）===")

    for filename, level in excel_files.items():
        file_path = TABLES_DIR / filename
        if not file_path.exists():
            print(f"  {filename}: 文件不存在，跳过")
            continue

        print(f"  正在读取 {filename}...")
        tables = read_excel_sheets(file_path)

        # 添加级别信息
        for table_name, table_info in tables.items():
            table_info['level'] = level
            all_tables[table_name] = table_info

        print(f"    提取了 {len(tables)} 个表")

    print(f"\n  总共提取了 {len(all_tables)} 个 S级表结构")
    return all_tables


def read_test_questions() -> Dict[str, List[Dict]]:
    """读取测试问题集合"""
    questions = {
        'simple': [],      # C级
        'join': [],        # B级
        'aggregate': [],   # A级
        'advanced': []     # S级
    }

    if not EXCEL_AVAILABLE:
        return questions

    test_file = TABLES_DIR / '测试问题集合.xlsx'
    if not test_file.exists():
        print("  测试问题集合.xlsx 文件不存在")
        return questions

    print("\n=== 读取测试问题集合 ===")

    try:
        import pandas as pd
        xl_file = pd.ExcelFile(test_file)

        for sheet_name in xl_file.sheet_names:
            df = pd.read_excel(test_file, sheet_name=sheet_name)

            # 打印列名用于调试
            print(f"  Sheet '{sheet_name}' 列名: {df.columns.tolist()}")

            # 检查列名 - 支持多种可能的列名
            col_mapping = {}
            for col in df.columns:
                col_lower = str(col).lower().strip()
                if '问题' in col and '分类' not in col:
                    col_mapping['question'] = col
                elif '问题分类' in col or '分类' in col:
                    col_mapping['category'] = col
                elif 'sql' in col.lower():
                    col_mapping['sql'] = col
                elif '答案' in col or '期望值' in col:
                    col_mapping['answer'] = col

            if 'question' not in col_mapping:
                print(f"  警告: Sheet '{sheet_name}' 缺少问题列")
                continue

            # 如果没有 sql 列，跳过
            if 'sql' not in col_mapping:
                print(f"  警告: Sheet '{sheet_name}' 缺少 SQL 列")
                continue

            for _, row in df.iterrows():
                question = str(row.get(col_mapping['question'], '')).strip()
                sql = str(row.get(col_mapping['sql'], '')).strip()

                # 跳过空值或 nan
                if not question or question.lower() == 'nan' or not sql or sql.lower() == 'nan':
                    continue

                # 获取分类/答案
                category = str(row.get(col_mapping.get('category', ''), '')).strip()
                answer = str(row.get(col_mapping.get('answer', ''), '')).strip()

                # 从问题分类推断级别
                level = 'C'
                if '第一批' in category:
                    level = 'C'
                elif '第二批' in category:
                    level = 'BC'
                elif '第三批' in category:
                    level = 'A'
                elif '第四批' in category or 'S' in category.upper():
                    level = 'S'

                q_item = {
                    'question': question,
                    'sql': sql,
                    'answer': answer if answer and answer.lower() != 'nan' else '',
                    'level': level
                }

                # 根据级别分类
                if level == 'C':
                    questions['simple'].append(q_item)
                elif level == 'BC' or level == 'B':
                    questions['join'].append(q_item)
                elif level == 'A':
                    questions['aggregate'].append(q_item)
                elif level == 'S':
                    questions['advanced'].append(q_item)
                else:
                    questions['simple'].append(q_item)

        total = sum(len(q) for q in questions.values())
        print(f"  读取了 {total} 个测试问题")
        for level, qs in questions.items():
            print(f"    {level}: {len(qs)} 个")

    except ImportError:
        print("  警告: pandas 未安装，无法读取测试问题")
        print("  请运行: pip install pandas openpyxl")
    except Exception as e:
        print(f"  读取测试问题时出错: {e}")
        import traceback
        traceback.print_exc()

    return questions


# ============================================================================
# 构建 DDL 知识库
# ============================================================================
def build_ddl_knowledge(tables: Dict[str, Any]):
    """构建 DDL 知识库"""
    print("\n=== 构建 DDL 知识库 ===")

    if not tables:
        print("  没有表结构数据")
        return False

    # 保存表结构 JSON
    output_file = DDL_DIR / 'tables.json'
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(tables, f, ensure_ascii=False, indent=2)

    print(f"  已生成表结构: {len(tables)} 个表")
    print(f"  保存到: {output_file}")

    # 提取表关系
    relationships = []
    for table_name, table_info in tables.items():
        for fk in table_info.get('foreign_keys', []):
            relationships.append({
                'from_table': table_name,
                'from_column': fk['column'],
                'to_table': fk['ref_table'],
                'to_column': fk['ref_column'],
                'relation': 'foreign_key'
            })

    # 保存表关系
    rel_file = DDL_DIR / 'relationships.json'
    with open(rel_file, 'w', encoding='utf-8') as f:
        json.dump(relationships, f, ensure_ascii=False, indent=2)

    print(f"  已生成表关系: {len(relationships)} 条")

    return True


# ============================================================================
# 构建业务描述知识库
# ============================================================================
def build_description_knowledge(tables: Dict[str, Any]):
    """构建业务描述知识库"""
    print("\n=== 构建业务描述知识库 ===")

    if not tables:
        print("  没有表结构数据")
        return False

    # 生成表描述
    tables_md = []
    tables_md.append("# 表业务描述\n\n")
    tables_md.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
    tables_md.append(f"总表数: {len(tables)}\n\n")
    tables_md.append("---\n\n")

    # 按级别分组
    by_level = {}
    for table_name, info in tables.items():
        level = info.get('level', 'Unknown')
        if level not in by_level:
            by_level[level] = []
        by_level[level].append((table_name, info))

    for level in ['S', 'A', 'BC', 'C']:
        if level not in by_level:
            continue
        tables_md.append(f"## {level}级表 ({len(by_level[level])}个)\n\n")

        for table_name, info in sorted(by_level[level]):
            tables_md.append(f"### {table_name}\n\n")
            tables_md.append(f"**业务概念**: {info.get('business_concept', '')}\n\n")
            tables_md.append(f"**说明**: {info.get('comment', '')}\n\n")
            if info.get('notes'):
                tables_md.append(f"**注意事项**: {info['notes']}\n\n")
            if info.get('department'):
                tables_md.append(f"**责任部门**: {info['department']}\n\n")

            # 字段列表
            columns = info.get('columns', [])
            if columns:
                tables_md.append("**字段列表**:\n\n")
                tables_md.append("| 字段名 | 中文名 | 类型 | 说明 |\n")
                tables_md.append("|--------|--------|------|------|\n")
                for col in columns[:20]:  # 限制显示前20个字段
                    tables_md.append(f"| {col['name']} | {col['name_cn']} | {col['type']} | {col['comment'][:50]} |\n")
                if len(columns) > 20:
                    tables_md.append(f"| ... | 还有 {len(columns)-20} 个字段 | | |\n")
                tables_md.append("\n")

    (DESC_DIR / 'tables.md').write_text(''.join(tables_md), encoding='utf-8')
    print(f"  已生成表描述: {len(tables)} 个表")

    # 生成字段索引
    fields_md = []
    fields_md.append("# 字段索引\n\n")
    fields_md.append("## 按表分类\n\n")

    for table_name, info in sorted(tables.items()):
        fields_md.append(f"### {table_name}\n\n")
        for col in info.get('columns', []):
            fields_md.append(f"- **{col['name']}** ({col['name_cn']}): {col['comment']}\n")
        fields_md.append("\n")

    (DESC_DIR / 'fields.md').write_text(''.join(fields_md), encoding='utf-8')
    print("  已生成字段索引")

    return True


# ============================================================================
# 构建查询示例知识库
# ============================================================================
def build_examples_knowledge(questions: Dict[str, List[Dict]]):
    """构建查询示例知识库"""
    print("\n=== 构建查询示例知识库 ===")

    total = sum(len(qs) for qs in questions.values())
    if total == 0:
        print("  没有测试问题数据，使用默认示例")

        # 使用默认示例
        default_examples = {
            'simple': """# 单表查询示例 (C级)

## 学校信息查询

**问题**: 我们学校的名字是什么？
```sql
SELECT NAME_ FROM HQ_JC_XX WHERE ROWNUM = 1;
```

**问题**: 学校有哪些校区？
```sql
SELECT NAME_ FROM HQ_JC_XQ ORDER BY NAME_;
```

## 学期查询

**问题**: 我们今年的学期范围是几号到几号？
```sql
SELECT SCHOOL_YEAR AS 学年, TERM_CODE AS 学期, TEACH_BEGIN_DATE AS 教学开始日期, TEACH_END_DATE AS 教学结束日期
FROM HQ_CODE_XNXQ ORDER BY BEGIN_DATE DESC FETCH FIRST 1 ROW ONLY;
```
""",
            'join': """# 多表联查示例 (B级)

## 学生-班级-专业联查

**问题**: 2024级学生属于哪些班级和专业？
```sql
SELECT s.STU_NO AS 学号, s.NAME_ AS 姓名, c.NAME_ AS 班级, m.NAME_ AS 专业, d.NAME_ AS 学院
FROM HQ_XS_STU s
LEFT JOIN HQ_CODE_CLASSES c ON s.CLASS_ID = c.NO_
LEFT JOIN HQ_CODE_MAJOR m ON s.MAJOR_CODE = m.CODE_
LEFT JOIN HQ_CODE_DEPT d ON s.DEPT_ID = d.ID
WHERE s.ENROLL_YEAR = '2024'
ORDER BY s.STU_NO;
```
""",
            'aggregate': """# 聚合查询示例 (A/S级)

## 计数统计

**问题**: 每个专业有多少学生？
```sql
SELECT m.NAME_ AS 专业名称, COUNT(s.STU_NO) AS 学生人数
FROM HQ_CODE_MAJOR m
LEFT JOIN HQ_XS_STU s ON m.CODE_ = s.MAJOR_CODE AND s.STU_STATE_CODE = '01'
GROUP BY m.NAME_
ORDER BY COUNT(s.STU_NO) DESC;
```
"""
        }

        (EX_DIR / 'simple.md').write_text(default_examples['simple'], encoding='utf-8')
        (EX_DIR / 'join.md').write_text(default_examples['join'], encoding='utf-8')
        (EX_DIR / 'aggregate.md').write_text(default_examples['aggregate'], encoding='utf-8')

        return True

    # 使用从 Excel 读取的问题
    for level, qs in questions.items():
        if not qs:
            continue

        content = []
        level_names = {
            'simple': '单表查询示例 (C级)',
            'join': '多表联查示例 (B级)',
            'aggregate': '聚合查询示例 (A级)',
            'advanced': '高级查询示例 (S级)'
        }

        content.append(f"# {level_names.get(level, level)}\n\n")
        content.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
        content.append(f"问题数量: {len(qs)}\n\n")
        content.append("---\n\n")

        for i, q in enumerate(qs[:100], 1):  # 限制每个级别最多100个
            content.append(f"## 示例 {i}\n\n")
            content.append(f"**问题**: {q['question']}\n\n")
            content.append(f"**难度**: {q['level']}\n\n")
            content.append("```sql\n")
            content.append(q['sql'])
            content.append("\n```\n\n")

            if q.get('answer'):
                content.append(f"**答案**: {q['answer']}\n\n")

            content.append("---\n\n")

        output_file = EX_DIR / f'{level}.md'
        output_file.write_text(''.join(content), encoding='utf-8')
        print(f"  已生成 {level} 级示例: {len(qs)} 个")

    return True


# ============================================================================
# 构建业务术语字典
# ============================================================================
def build_glossary(tables: Dict[str, Any]):
    """构建业务术语字典"""
    print("\n=== 构建业务术语字典 ===")

    if not tables:
        print("  没有表结构数据")
        return False

    glossary = []
    glossary.append("# 业务术语知识库\n\n")
    glossary.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
    glossary.append(f"总表数: {len(tables)}\n\n")
    glossary.append("---\n\n")

    # 表名映射
    glossary.append("## 实际表名映射\n\n")
    glossary.append("| 业务概念 | 实际表名 | 说明 | 级别 |\n")
    glossary.append("|----------|----------|------|------|\n")

    for table_name, info in sorted(tables.items()):
        concept = info.get('business_concept', '')
        desc = info.get('comment', '')
        level = info.get('level', '')
        glossary.append(f"| {concept} | {table_name} | {desc} | {level} |\n")

    glossary.append("\n---\n\n")

    # 常见查询意图
    glossary.append("## 常见查询意图映射\n\n")
    glossary.append("| 用户问题 | 语义 | SQL 关键字 |\n")
    glossary.append("|----------|------|------------|\n")
    glossary.append("| 有多少 | 计数 | COUNT(*) |\n")
    glossary.append("| 数量 | 计数 | COUNT(*) |\n")
    glossary.append("| 总数 | 计数 | COUNT(*) |\n")
    glossary.append("| 平均 | 平均值 | AVG() |\n")
    glossary.append("| 比例 | 百分比 | CASE WHEN + COUNT/SUM |\n")
    glossary.append("| 排名 | 排序 | ORDER BY + RANK() |\n")
    glossary.append("| 是什么 | 单值查询 | SELECT ... WHERE ... ROWNUM = 1 |\n")
    glossary.append("| 有哪些 | 列表查询 | SELECT ... |\n")
    glossary.append("| 哪些 | 筛选查询 | SELECT ... WHERE |\n")

    # 保存到 .claude/knowledge/ 目录
    output_dir = Path(__file__).parent.parent / '.claude' / 'knowledge'
    output_dir.mkdir(parents=True, exist_ok=True)

    output_file = output_dir / 'business-glossary-auto.md'
    output_file.write_text(''.join(glossary), encoding='utf-8')

    print(f"  已生成业务术语字典: {len(tables)} 个表")
    print(f"  保存到: {output_file}")

    return True


# ============================================================================
# 主函数
# ============================================================================
def main():
    """主函数"""
    print("=" * 60)
    print("RAGFlow 增强版知识库构建脚本")
    print("=" * 60)
    print(f"输出目录: {OUTPUT_DIR}")
    print(f"Tables 目录: {TABLES_DIR}")
    print()

    # 确保目录存在
    DDL_DIR.mkdir(parents=True, exist_ok=True)
    DESC_DIR.mkdir(parents=True, exist_ok=True)
    EX_DIR.mkdir(parents=True, exist_ok=True)

    # 检查依赖
    if not EXCEL_AVAILABLE:
        print("=" * 60)
        print("警告: openpyxl 未安装！")
        print("请运行: pip install openpyxl")
        print("=" * 60)
        return

    # 读取 Excel 数据
    tables = read_all_excel_tables()
    questions = read_test_questions()

    # 构建知识库
    success = True
    if not build_ddl_knowledge(tables):
        success = False
    if not build_description_knowledge(tables):
        success = False
    if not build_examples_knowledge(questions):
        success = False
    if not build_glossary(tables):
        success = False

    print()
    print("=" * 60)
    if success:
        print("知识库构建完成！")
        print("=" * 60)
        print()
        print("生成文件:")
        print(f"  - {DDL_DIR}/tables.json: 表结构数据")
        print(f"  - {DDL_DIR}/relationships.json: 表关系数据")
        print(f"  - {DESC_DIR}/tables.md: 表业务描述")
        print(f"  - {DESC_DIR}/fields.md: 字段索引")
        print(f"  - {EX_DIR}/*.md: 查询示例")
        print()
        print("下一步:")
        print("1. 重启 RAG MCP 服务")
        print("2. 测试查询")
    else:
        print("部分知识库构建失败，请检查错误信息")

if __name__ == "__main__":
    main()
